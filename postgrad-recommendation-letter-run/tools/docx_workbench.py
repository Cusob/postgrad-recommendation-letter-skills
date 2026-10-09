#!/usr/bin/env python3
"""Small, conservative DOCX inspection and editing workbench.

The workbench handles mechanical operations only. It never invents letter
content and it never rebuilds a document from paragraph strings.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

try:
    from docx import Document
except ImportError as exc:  # pragma: no cover - environment-dependent
    print("python-docx is required: %s" % exc, file=sys.stderr)
    raise SystemExit(2)


TOOL_VERSION = "3.0"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W_NS, "wp": WP_NS}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def scalar(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "inches"):
        return round(float(value.inches), 6)
    if hasattr(value, "pt"):
        return round(float(value.pt), 4)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def rgb_value(color: Any) -> Any:
    if color is None:
        return None
    try:
        return str(color.rgb) if color.rgb is not None else None
    except (AttributeError, ValueError):
        return None


def run_style(run: Any) -> Dict[str, Any]:
    font = run.font
    return {
        "font_name": font.name,
        "font_size_pt": scalar(font.size),
        "bold": font.bold,
        "italic": font.italic,
        "underline": scalar(font.underline),
        "strike": font.strike,
        "color": rgb_value(font.color),
        "highlight": scalar(font.highlight_color),
        "style": run.style.name if run.style is not None else None,
    }


def run_signature(run_data: Dict[str, Any]) -> str:
    return json.dumps(run_data, ensure_ascii=False, sort_keys=True)


def paragraph_properties(paragraph: Any) -> Dict[str, Any]:
    fmt = paragraph.paragraph_format
    return {
        "style": paragraph.style.name if paragraph.style is not None else None,
        "alignment": scalar(paragraph.alignment),
        "left_indent": scalar(fmt.left_indent),
        "right_indent": scalar(fmt.right_indent),
        "first_line_indent": scalar(fmt.first_line_indent),
        "space_before": scalar(fmt.space_before),
        "space_after": scalar(fmt.space_after),
        "line_spacing": scalar(fmt.line_spacing),
        "line_spacing_rule": scalar(fmt.line_spacing_rule),
        "keep_together": fmt.keep_together,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "widow_control": fmt.widow_control,
    }


def style_record(style: Any) -> Dict[str, Any]:
    font = style.font
    record = {
        "name": style.name,
        "style_id": style.style_id,
        "type": scalar(style.type),
        "base_style": style.base_style.name if style.base_style is not None else None,
        "font": {
            "name": font.name,
            "size_pt": scalar(font.size),
            "bold": font.bold,
            "italic": font.italic,
            "underline": scalar(font.underline),
            "color": rgb_value(font.color),
        },
    }
    try:
        fmt = style.paragraph_format
        record["paragraph"] = {
            "alignment": scalar(fmt.alignment),
            "left_indent": scalar(fmt.left_indent),
            "right_indent": scalar(fmt.right_indent),
            "first_line_indent": scalar(fmt.first_line_indent),
            "space_before": scalar(fmt.space_before),
            "space_after": scalar(fmt.space_after),
            "line_spacing": scalar(fmt.line_spacing),
            "line_spacing_rule": scalar(fmt.line_spacing_rule),
        }
    except (AttributeError, ValueError):
        record["paragraph"] = None
    return record


def paragraph_record(scope: str, paragraph: Any) -> Dict[str, Any]:
    runs = []
    cursor = 0
    for run in paragraph.runs:
        text = run.text or ""
        style = run_style(run)
        runs.append({
            "text": text,
            "start": cursor,
            "end": cursor + len(text),
            "style": style,
            "style_signature": run_signature(style),
        })
        cursor += len(text)
    element = paragraph._p
    return {
        "scope": scope,
        "text": paragraph.text,
        "properties": paragraph_properties(paragraph),
        "runs": runs,
        "run_count": len(runs),
        "nonempty_style_signatures": sorted({
            item["style_signature"] for item in runs if item["text"]
        }),
        "placeholders": sorted(set(re.findall(r"\{\{[^{}]+\}\}", paragraph.text))),
        "hyperlink_count": len(element.findall(".//w:hyperlink", NS)),
        "drawing_count": len(element.findall(".//w:drawing", NS)),
        "textbox_count": len(element.findall(".//w:txbxContent", NS)),
    }


def iter_cell_paragraphs(cell: Any, prefix: str) -> Iterable[Tuple[str, Any]]:
    for index, paragraph in enumerate(cell.paragraphs):
        yield "%s/p:%d" % (prefix, index), paragraph
    for table_index, table in enumerate(cell.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, nested_cell in enumerate(row.cells):
                nested_prefix = "%s/table:%d/r:%d/c:%d" % (
                    prefix,
                    table_index,
                    row_index,
                    cell_index,
                )
                yield from iter_cell_paragraphs(nested_cell, nested_prefix)


def iter_table_paragraphs(table: Any, prefix: str) -> Iterable[Tuple[str, Any]]:
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell_prefix = "%s/r:%d/c:%d" % (prefix, row_index, cell_index)
            yield from iter_cell_paragraphs(cell, cell_prefix)


def iter_paragraphs(document: Any) -> Iterable[Tuple[str, Any]]:
    for index, paragraph in enumerate(document.paragraphs):
        yield "body/p:%d" % index, paragraph
    for table_index, table in enumerate(document.tables):
        yield from iter_table_paragraphs(table, "table:%d" % table_index)
    for section_index, section in enumerate(document.sections):
        for label, container in (
            ("header", section.header),
            ("first_header", section.first_page_header),
            ("even_header", section.even_page_header),
            ("footer", section.footer),
            ("first_footer", section.first_page_footer),
            ("even_footer", section.even_page_footer),
        ):
            container_prefix = "section:%d/%s" % (section_index, label)
            for paragraph_index, paragraph in enumerate(container.paragraphs):
                yield "%s/p:%d" % (
                    container_prefix,
                    paragraph_index,
                ), paragraph
            for table_index, table in enumerate(container.tables):
                yield from iter_table_paragraphs(
                    table,
                    "%s/table:%d" % (container_prefix, table_index),
                )


def count_tables(document: Any) -> int:
    count = 0

    def nested_count(table: Any) -> int:
        total = 1
        for row in table.rows:
            for cell in row.cells:
                total += sum(nested_count(nested) for nested in cell.tables)
        return total

    for table in document.tables:
        count += nested_count(table)
    for section in document.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            for table in container.tables:
                count += nested_count(table)
    return count


def section_record(section: Any) -> Dict[str, Any]:
    containers = {}
    for label, container in (
        ("header", section.header),
        ("first_header", section.first_page_header),
        ("even_header", section.even_page_header),
        ("footer", section.footer),
        ("first_footer", section.first_page_footer),
        ("even_footer", section.even_page_footer),
    ):
        part = container.part
        containers[label] = {
            "part_name": str(part.partname) if part is not None else None,
            "linked_to_previous": container.is_linked_to_previous,
            "paragraphs": len(container.paragraphs),
            "tables": len(container.tables),
        }
    return {
        "top_margin": scalar(section.top_margin),
        "bottom_margin": scalar(section.bottom_margin),
        "left_margin": scalar(section.left_margin),
        "right_margin": scalar(section.right_margin),
        "header_distance": scalar(section.header_distance),
        "footer_distance": scalar(section.footer_distance),
        "orientation": scalar(section.orientation),
        "page_width": scalar(section.page_width),
        "page_height": scalar(section.page_height),
        "containers": containers,
    }


def relationship_summary(document: Any) -> Dict[str, int]:
    counts = {"image": 0, "hyperlink": 0, "ole_object": 0}
    for rel in document.part.rels.values():
        target = str(rel.target_ref or "")
        if "image" in rel.reltype:
            counts["image"] += 1
        elif "hyperlink" in rel.reltype:
            counts["hyperlink"] += 1
        elif "oleObject" in target or "oleObject" in rel.reltype:
            counts["ole_object"] += 1
    return counts


def document_snapshot(path: Path) -> Dict[str, Any]:
    document = Document(str(path))
    paragraphs = [
        paragraph_record(scope, paragraph)
        for scope, paragraph in iter_paragraphs(document)
    ]
    used_style_names = {item["properties"]["style"] for item in paragraphs if item["properties"]["style"]}
    used_style_names.update(
        run["style"]["style"]
        for item in paragraphs
        for run in item["runs"]
        if run["style"].get("style")
    )
    all_styles = {style.name: style for style in document.styles}
    pending = list(used_style_names)
    while pending:
        name = pending.pop()
        style = all_styles.get(name)
        if style is not None and style.base_style is not None and style.base_style.name not in used_style_names:
            used_style_names.add(style.base_style.name)
            pending.append(style.base_style.name)
    styles = [style_record(all_styles[name]) for name in sorted(used_style_names) if name in all_styles]
    with zipfile.ZipFile(str(path)) as archive:
        names = archive.namelist()
        document_xml = archive.read("word/document.xml").decode("utf-8", "ignore")
        has_comments = any(name.startswith("word/comments") for name in names)
    return {
        "tool_version": TOOL_VERSION,
        "path": str(path),
        "sha256": sha256_file(path),
        "counts": {
            "paragraphs": len(paragraphs),
            "body_paragraphs": len(document.paragraphs),
            "tables": count_tables(document),
            "sections": len(document.sections),
            "inline_shapes": len(document.inline_shapes),
            "paragraph_drawings": sum(p["drawing_count"] for p in paragraphs),
            "textboxes": sum(p["textbox_count"] for p in paragraphs),
            "hyperlinks": sum(p["hyperlink_count"] for p in paragraphs),
        },
        "sections": [section_record(section) for section in document.sections],
        "styles": styles,
        "relationships": relationship_summary(document),
        "has_comments": has_comments,
        "has_document_drawing": "<w:drawing" in document_xml,
        "paragraphs": paragraphs,
    }


def write_json(data: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def locate_runs(paragraph: Any) -> List[Tuple[Any, int, int]]:
    result = []
    cursor = 0
    for run in paragraph.runs:
        text = run.text or ""
        result.append((run, cursor, cursor + len(text)))
        cursor += len(text)
    return result


def replace_range(paragraph: Any, start: int, end: int, replacement: str) -> Dict[str, Any]:
    locations = locate_runs(paragraph)
    if not locations:
        return {"ok": False, "reason": "no_editable_runs"}
    start_run = None
    end_run = None
    for run, run_start, run_end in locations:
        if start_run is None and start < run_end:
            start_run = (run, run_start, run_end)
        if end <= run_end:
            end_run = (run, run_start, run_end)
            break
    if start_run is None or end_run is None:
        return {"ok": False, "reason": "range_not_in_runs"}
    first, first_start, _ = start_run
    last, last_start, _ = end_run
    first_text = first.text or ""
    if first is last:
        local_start = start - first_start
        local_end = end - first_start
        first.text = first_text[:local_start] + replacement + first_text[local_end:]
        return {"ok": True, "run_span": 1}
    first_text = first.text or ""
    last_text = last.text or ""
    first_local_start = start - first_start
    last_local_end = end - last_start
    first.text = first_text[:first_local_start] + replacement
    last.text = last_text[last_local_end:]
    seen_first = False
    for run, _, _ in locations:
        if run is first:
            seen_first = True
            continue
        if run is last:
            break
        if seen_first:
            run.text = ""
    return {"ok": True, "run_span": sum(1 for run, _, _ in locations if run is not first and run is not last) + 2}


def all_text_replacements(document: Any, replacements: Dict[str, str]) -> Dict[str, Any]:
    changes = []
    unresolved = []
    ordered = sorted(
        ((str(old), str(new)) for old, new in replacements.items() if str(old)),
        key=lambda item: len(item[0]),
        reverse=True,
    )
    for scope, paragraph in iter_paragraphs(document):
        editable_text = "".join(run.text or "" for run in paragraph.runs)
        for old, new in ordered:
            while old in editable_text:
                full_text = editable_text
                start = full_text.find(old)
                outcome = replace_range(paragraph, start, start + len(old), new)
                if not outcome.get("ok"):
                    unresolved.append({"scope": scope, "old": old, "reason": outcome["reason"]})
                    break
                changes.append({
                    "scope": scope,
                    "old": old,
                    "new": new,
                    "run_span": outcome.get("run_span"),
                })
                editable_text = "".join(run.text or "" for run in paragraph.runs)
    return {"changes": changes, "unresolved": unresolved}


def paragraph_replacements(document: Any, operations: List[Dict[str, Any]]) -> Dict[str, Any]:
    paragraphs = list(iter_paragraphs(document))
    changes = []
    unresolved = []
    for operation in operations:
        scope = str(operation.get("scope", "body/p:%d" % operation.get("index", -1)))
        new_text = str(operation.get("text", ""))
        matches = [(key, paragraph) for key, paragraph in paragraphs if key == scope]
        if not matches:
            unresolved.append({"scope": scope, "reason": "scope_not_found"})
            continue
        _, paragraph = matches[0]
        if paragraph._p.findall(".//w:hyperlink", NS):
            unresolved.append({"scope": scope, "reason": "paragraph_contains_hyperlink"})
            continue
        old_text = "".join(run.text or "" for run in paragraph.runs)
        if not old_text:
            unresolved.append({"scope": scope, "reason": "empty_paragraph"})
            continue
        outcome = replace_range(paragraph, 0, len(old_text), new_text)
        if outcome.get("ok"):
            changes.append({"scope": scope, "old": old_text, "new": new_text, **outcome})
        else:
            unresolved.append({"scope": scope, "reason": outcome["reason"]})
    return {"changes": changes, "unresolved": unresolved}


def compare_snapshots(baseline: Dict[str, Any], current: Dict[str, Any]) -> Dict[str, Any]:
    critical = []
    warnings = []
    info = []

    for key in ("paragraphs", "tables", "sections", "inline_shapes", "textboxes"):
        before = baseline.get("counts", {}).get(key)
        after = current.get("counts", {}).get(key)
        if before != after:
            critical.append({"kind": "count_changed", "field": key, "before": before, "after": after})

    before_sections = baseline.get("sections", [])
    after_sections = current.get("sections", [])
    if len(before_sections) == len(after_sections):
        for index, (before, after) in enumerate(zip(before_sections, after_sections)):
            if before != after:
                critical.append({"kind": "section_properties_changed", "index": index, "before": before, "after": after})

    if baseline.get("styles", []) != current.get("styles", []):
        critical.append({
            "kind": "used_style_definitions_changed",
            "before": baseline.get("styles", []),
            "after": current.get("styles", []),
        })

    before_by_scope = {item["scope"]: item for item in baseline.get("paragraphs", [])}
    after_by_scope = {item["scope"]: item for item in current.get("paragraphs", [])}
    for scope in sorted(set(before_by_scope) | set(after_by_scope)):
        before = before_by_scope.get(scope)
        after = after_by_scope.get(scope)
        if before is None or after is None:
            critical.append({"kind": "paragraph_scope_changed", "scope": scope})
            continue
        if before["properties"] != after["properties"]:
            critical.append({"kind": "paragraph_properties_changed", "scope": scope})
        if before.get("hyperlink_count") != after.get("hyperlink_count"):
            critical.append({"kind": "hyperlink_count_changed", "scope": scope})
        if before.get("drawing_count") != after.get("drawing_count"):
            critical.append({"kind": "drawing_count_changed", "scope": scope})
        if before.get("textbox_count") != after.get("textbox_count"):
            critical.append({"kind": "textbox_count_changed", "scope": scope})
        if before.get("text") != after.get("text"):
            info.append({"kind": "text_changed", "scope": scope})
        if before.get("run_count") != after.get("run_count"):
            warnings.append({
                "kind": "run_count_changed",
                "scope": scope,
                "before": before.get("run_count"),
                "after": after.get("run_count"),
            })
        before_styles = set(before.get("nonempty_style_signatures", []))
        after_styles = set(after.get("nonempty_style_signatures", []))
        if len(before_styles) > 1 and len(after_styles) <= 1 and after.get("text"):
            critical.append({"kind": "run_style_collapse", "scope": scope})
        if not after_styles.issubset(before_styles) and after_styles:
            warnings.append({"kind": "new_run_style", "scope": scope})

    for key in ("image", "hyperlink", "ole_object"):
        before = baseline.get("relationships", {}).get(key)
        after = current.get("relationships", {}).get(key)
        if before != after:
            critical.append({"kind": "relationship_count_changed", "field": key, "before": before, "after": after})
    if baseline.get("has_comments") != current.get("has_comments"):
        critical.append({"kind": "comments_presence_changed"})
    return {
        "status": "critical_failure" if critical else "pass_with_warnings" if warnings else "pass",
        "critical": critical,
        "warnings": warnings,
        "info": info,
        "baseline_sha256": baseline.get("sha256"),
        "generated_sha256": current.get("sha256"),
    }


def command_inspect(args: argparse.Namespace) -> int:
    data = document_snapshot(Path(args.docx).resolve())
    if args.output:
        write_json(data, Path(args.output).resolve())
    else:
        print(json.dumps(data, ensure_ascii=False, indent=2))
    return 0


def command_snapshot(args: argparse.Namespace) -> int:
    data = document_snapshot(Path(args.docx).resolve())
    output = Path(args.output).resolve()
    write_json(data, output)
    print("snapshot written: %s" % output)
    return 0


def load_json(path: str) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def command_replace(args: argparse.Namespace) -> int:
    source = Path(args.docx).resolve()
    output = Path(args.output).resolve()
    document = Document(str(source))
    result = {"changes": [], "unresolved": []}
    if args.map:
        mapping = load_json(args.map)
        if not isinstance(mapping, dict):
            raise ValueError("replacement map must be a JSON object")
        result = all_text_replacements(document, mapping)
    if args.paragraphs:
        operations = load_json(args.paragraphs)
        if not isinstance(operations, list):
            raise ValueError("paragraph operations must be a JSON list")
        paragraph_result = paragraph_replacements(document, operations)
        result["changes"].extend(paragraph_result["changes"])
        result["unresolved"].extend(paragraph_result["unresolved"])
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    result["source"] = str(source)
    result["output"] = str(output)
    if args.report:
        write_json(result, Path(args.report).resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 2 if result["unresolved"] else 0


def command_compare(args: argparse.Namespace) -> int:
    baseline = load_json(args.snapshot)
    current = document_snapshot(Path(args.docx).resolve())
    report = compare_snapshots(baseline, current)
    if args.output:
        write_json(report, Path(args.output).resolve())
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 2 if report["critical"] else 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Conservative DOCX inspection and comparison tool")
    sub = parser.add_subparsers(dest="command", required=True)

    inspect = sub.add_parser("inspect")
    inspect.add_argument("docx")
    inspect.add_argument("--output")
    inspect.set_defaults(func=command_inspect)

    snapshot = sub.add_parser("snapshot")
    snapshot.add_argument("docx")
    snapshot.add_argument("--output", required=True)
    snapshot.set_defaults(func=command_snapshot)

    replace = sub.add_parser("replace")
    replace.add_argument("docx")
    replace.add_argument("--output", required=True)
    replace.add_argument("--map", help="JSON object of exact text replacements")
    replace.add_argument("--paragraphs", help="JSON list of scoped paragraph rewrite operations")
    replace.add_argument("--report")
    replace.set_defaults(func=command_replace)

    compare = sub.add_parser("compare")
    compare.add_argument("snapshot")
    compare.add_argument("docx")
    compare.add_argument("--output")
    compare.set_defaults(func=command_compare)
    return parser


def main(argv: List[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return args.func(args)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        print("docx_workbench error: %s" % exc, file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
